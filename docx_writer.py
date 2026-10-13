"""Zet de vertaling (in eenvoudige Markdown) om naar een net Word-document.

Gebruikt voor PDF's, scans, foto's en tekstbestanden. De vertaler geeft de
structuur aan (koppen, vet/cursief, tabellen, opsommingen, nieuwe pagina's)
en hier wordt dat een echt .docx-bestand met dezelfde opbouw.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

from docx_layout import W_P, W_R, _w, set_direction

# Regel waarmee de vertaler aangeeft dat in het origineel een nieuwe pagina begint.
PAGE_MARKER = "<<<PAGE>>>"

_HEADER_LABELS = {
    "nederlands": "Vertaling van",
    "english": "Translation of",
    "arabic": "ترجمة لـ",
}

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*•]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^\(?(\d{1,3}|[a-zA-Z])[.)]\s+")
_RULE_RE = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")
_TABLE_SEPARATOR_RE =re.compile(r"^:?-{2,}:?$")
_INLINE_RE = re.compile(r"\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])")
_BR_RE = re.compile(r"<br\s*/?>", re.IGNORECASE)


def _strip_code_fence(text: str) -> str:
    lines = text.strip().splitlines()
    if lines and lines[0].startswith("```"):
        lines = lines[1:]
        if lines and lines[-1].strip().startswith("```"):
            lines = lines[:-1]
    return "\n".join(lines)


def _add_inline(paragraph, text: str):
    """Voeg tekst toe met **vet**, *cursief* en ***beide***."""
    pos = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        both, bold, italic = match.groups()
        run = paragraph.add_run(both or bold or italic)
        run.bold = bool(both or bold)
        run.italic = bool(both or italic)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def _add_table(document, rows: list[list[str]]):
    rows = [r for r in rows if not all(_TABLE_SEPARATOR_RE.match(c) for c in r if c)]
    if not rows:
        return
    columns = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(columns):
            cell = table.cell(r_idx, c_idx)
            parts = _BR_RE.split(row[c_idx]) if c_idx < len(row) else [""]
            _add_inline(cell.paragraphs[0], parts[0])
            for extra in parts[1:]:
                _add_inline(cell.add_paragraph(), extra)
    # Lege alinea na de tabel, zodat twee tabellen niet aan elkaar plakken.
    document.add_paragraph()


def _setup_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(4)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    for section in document.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)


def _make_rtl(document):
    """Alles van rechts naar links: alinea's, runs en de kolomvolgorde van tabellen."""
    roots = [document.element.body] + [s.header._element for s in document.sections]
    for root in roots:
        for p in root.iter(W_P):
            set_direction(p, list(p.iter(W_R)), True)
    for tbl_pr in document.element.body.iter(_w("tblPr")):
        if tbl_pr.find(_w("bidiVisual")) is None:
            # bidiVisual hoort direct na tblStyle/tblpPr/tblOverlap.
            idx = sum(1 for c in tbl_pr if c.tag in (_w("tblStyle"), _w("tblpPr"), _w("tblOverlap")))
            tbl_pr.insert(idx, etree.Element(_w("bidiVisual")))


def markdown_to_text(markdown: str) -> str:
    """Maak van de vertaling (eenvoudige Markdown) gewone tekst zonder opmaaktekens.

    Koppen, vet/cursief en opsommingen worden gewone regels, tabelrijen worden
    cellen gescheiden door een tab, en een nieuwe pagina wordt een witregel extra.
    """
    out = []
    for raw in _strip_code_fence(markdown).splitlines():
        line = raw.strip()
        if line == PAGE_MARKER:
            out += ["", ""]
            continue
        if _RULE_RE.match(line):
            continue
        if line.startswith("|"):
            cells = _split_row(line)
            if cells and all(_TABLE_SEPARATOR_RE.match(c) for c in cells if c):
                continue
            line = "\t".join(_BR_RE.sub(" ", c) for c in cells)
        else:
            heading = _HEADING_RE.match(line)
            if heading:
                line = heading.group(2)
            line = _BR_RE.sub("\n", line)
        line = _INLINE_RE.sub(lambda m: m.group(1) or m.group(2) or m.group(3), line)
        out.append(line)
    text = "\n".join(out)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip() + "\n"


def build_docx(markdown: str, dest_path: Path, source_name: str, target_language: str, rtl: bool):
    document = Document()
    _setup_styles(document)

    label = _HEADER_LABELS.get(target_language.strip().lower(), "Vertaling van")
    header = document.sections[0].header.paragraphs[0]
    header_run = header.add_run(f"{label}: {source_name}")
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    lines = _strip_code_fence(markdown).splitlines()
    has_content = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if line == PAGE_MARKER:
            if has_content:
                document.add_page_break()
            i += 1
            continue

        # Lege regels en horizontale lijnen ("---") overslaan.
        if not line or _RULE_RE.match(line):
            i += 1
            continue

        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            _add_table(document, rows)
            has_content = True
            continue

        heading = _HEADING_RE.match(line)
        bullet = _BULLET_RE.match(line)
        numbered = _NUMBERED_RE.match(line)
        if heading:
            level = min(len(heading.group(1)), 3)
            paragraph = document.add_heading(level=level)
            _add_inline(paragraph, heading.group(2))
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, bullet.group(1))
        elif numbered:
            # Nummering letterlijk overnemen ("3.", "b)") i.p.v. Word te laten hernummeren.
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            _add_inline(paragraph, line)
        else:
            paragraph = document.add_paragraph()
            parts = _BR_RE.split(line)
            _add_inline(paragraph, parts[0])
            for extra in parts[1:]:
                paragraph.add_run().add_break()
                _add_inline(paragraph, extra)
        has_content = True
        i += 1

    if rtl:
        _make_rtl(document)
    document.save(str(dest_path))
